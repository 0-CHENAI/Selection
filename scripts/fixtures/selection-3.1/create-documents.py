"""Deterministic native-format acceptance inputs; not a pre-authored plan or model output."""
from pathlib import Path
import sys
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from pptx import Presentation
root = Path(sys.argv[1]); root.mkdir(parents=True, exist_ok=True)
pdf = canvas.Canvas(str(root / 'cost.pdf'))
for page in range(1, 21):
    pdf.drawString(60, 760, f'TEST FIXTURE ONLY. Physical page {page}.')
    if page == 17:
        pdf.drawString(60, 730, 'Plan A two-year cost: CNY 1000000 (100 wan yuan).')
        pdf.drawString(60, 700, 'Plan B basis is pending; do not compare directly.')
        pdf.drawString(60, 670, 'Risk information has one unresolved gap; details are unavailable.')
    else: pdf.drawString(60, 730, 'Background only; no cost or risk claim on this page.')
    pdf.showPage()
pdf.save()
doc = Document(); doc.add_heading('Test scope', level=1); doc.add_paragraph('Fixtures only; no market fact.')
for chapter in range(1, 101):
    doc.add_heading(f'Background {chapter}', level=1)
    doc.add_paragraph('This background paragraph contains no cost conclusion. ' * 20)
doc.add_heading('Cost source', level=1); doc.add_paragraph('Plan A two-year cost: CNY 1000000.')
doc.save(root / 'cost.docx')
book = Workbook(); sheet = book.active; sheet.title = 'Costs'
sheet['A1'] = 'TEST FIXTURE ONLY'; sheet['A2'] = 'A two-year cost CNY'; sheet['B2'] = 1000000
sheet['A3'] = 'B basis pending; no direct comparison'; sheet['A4'] = 'Risk gap exists; details unavailable'
book.save(root / 'cost.xlsx')
deck = Presentation()
for title, body in [('Test scope', 'Fixtures only; no market fact.'), ('Cost source', 'Plan A two-year cost: CNY 1000000.')]:
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
deck.save(root / 'cost.pptx')
